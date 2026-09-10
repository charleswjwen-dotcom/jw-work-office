"""读回 modified.pptx，校验改写是否生效 —— 对应验收标准 1「保存后可正常打开且改动存在」。"""
import sys
import json
from pptx import Presentation


def verify(path: str) -> dict:
    prs = Presentation(path)
    texts = []
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append((i, shape.text_frame.text.strip()))

    checks = {
        "opened_ok": True,
        "slide_count": len(prs.slides._sldIdLst),
        "has_modified_title": any("已修改-" in t for _, t in texts),
        "has_poc_prefix": any("[PoC 改写]" in t for _, t in texts),
        "texts": texts,
    }
    checks["pass"] = checks["has_modified_title"] and checks["has_poc_prefix"]
    return checks


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "modified.pptx"
    print(json.dumps(verify(path), ensure_ascii=False, indent=2))
