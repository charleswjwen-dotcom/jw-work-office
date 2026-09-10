"""生成一份测试用 .pptx，含文本框、标题、图片、形状，供 PoC 改写验证使用。"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image


def make_sample_image(path):
    img = Image.new("RGB", (400, 300), color=(70, 130, 180))
    img.save(path)


def build(out_path, img_path):
    make_sample_image(img_path)
    prs = Presentation()

    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "原始标题"
    slide1.placeholders[1].text = "原始副标题内容"

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tb.text_frame.text = "占位文本-待替换"

    box = slide2.shapes.add_shape(
        1, Inches(1), Inches(3), Inches(2), Inches(1)
    )
    box.text_frame.text = "形状文字"
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xC0, 0xC0, 0xC0)

    slide2.shapes.add_picture(img_path, Inches(4), Inches(3), Inches(2), Inches(1.5))

    prs.save(out_path)
    print(f"[build] saved sample -> {out_path}")
    print(f"[build] slides={len(prs.slides._sldIdLst)}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "sample.pptx"
    img = sys.argv[2] if len(sys.argv) > 2 else "orig.png"
    build(out, img)
