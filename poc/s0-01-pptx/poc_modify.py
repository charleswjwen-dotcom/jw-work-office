"""
T-S0-01 PoC：验证 python-pptx 读取/修改/保存真实 .pptx 的全链路。
验收标准 1：能读取并原地修改文本、图片、形状属性，保存后文件结构完整。
"""
import sys
import json
import time
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from PIL import Image


def replace_image_blob(slide, new_color: tuple):
    from pptx.util import Inches
    import io
    for shape in slide.shapes:
        if shape.shape_type == 13:
            img = Image.new("RGB", (400, 300), color=new_color)
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            buf.seek(0)
            pic = shape._element
            blip = pic.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
            if blip is not None:
                rId = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                image_part = slide.part.related_part(rId)
                image_part._blob = buf.read()
                return True
    return False


def modify(in_path: str, out_path: str) -> dict:
    start = time.time()
    prs = Presentation(in_path)
    results = {"slides_modified": [], "image_replaced": False}

    slide0 = prs.slides[0]
    for shape in slide0.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if "原始" in run.text:
                        run.text = run.text.replace("原始", "已修改-")
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xE0, 0x40, 0x40)
            results["slides_modified"].append(0)

    slide1 = prs.slides[1]
    for shape in slide1.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if "占位" in run.text or "形状" in run.text:
                        run.text = "[PoC 改写] " + run.text
            results["slides_modified"].append(1)

    results["image_replaced"] = replace_image_blob(slide1, (220, 80, 60))

    prs.save(out_path)
    results["elapsed_ms"] = round((time.time() - start) * 1000, 1)
    results["out_size_kb"] = round(Path(out_path).stat().st_size / 1024, 1)
    return results


if __name__ == "__main__":
    in_path  = sys.argv[1] if len(sys.argv) > 1 else "sample.pptx"
    out_path = sys.argv[2] if len(sys.argv) > 2 else "modified.pptx"
    r = modify(in_path, out_path)
    print(json.dumps(r, ensure_ascii=False, indent=2))
