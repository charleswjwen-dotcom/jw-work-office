import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(__file__), "samples")
os.makedirs(OUT, exist_ok=True)

LOREM = (
    "人工智能正在深刻改变现代办公的方式，从文档自动处理到智能排版，"
    "效率提升已成为企业数字化转型的核心命题。"
)

def _add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)

def make_simple():
    doc = Document()
    doc.core_properties.title = "简单测试文档"
    for page in range(1, 4):
        doc.add_heading(f"第 {page} 页标题", level=1)
        for i in range(1, 9):
            p = doc.add_paragraph(f"[P{page}-{i}] {LOREM}")
            p.paragraph_format.space_after = Pt(6)
        if page < 3:
            _add_page_break(doc)
    doc.save(os.path.join(OUT, "simple.docx"))
    print("生成: simple.docx  (3页, 纯段落)")

def make_medium():
    doc = Document()
    doc.core_properties.title = "中等复杂度测试文档"
    for page in range(1, 4):
        doc.add_heading(f"第 {page} 章：结构化内容", level=1)
        doc.add_heading(f"1.{page} 概述", level=2)
        for i in range(1, 5):
            p = doc.add_paragraph(f"[M{page}-{i}] {LOREM}")
            p.paragraph_format.space_after = Pt(4)
        doc.add_heading(f"1.{page}.1 数据表格", level=3)
        tbl = doc.add_table(rows=4, cols=3)
        tbl.style = "Table Grid"
        headers = ["编号", "名称", "备注"]
        for j, h in enumerate(headers):
            tbl.rows[0].cells[j].text = h
        for r in range(1, 4):
            for c in range(3):
                tbl.rows[r].cells[c].text = f"R{r}C{c+1}-P{page}"
        if page < 3:
            _add_page_break(doc)
    doc.save(os.path.join(OUT, "medium.docx"))
    print("生成: medium.docx  (3页, 标题+表格)")

def make_complex():
    doc = Document()
    doc.core_properties.title = "复杂测试文档"

    section = doc.sections[0]
    section.header_distance = Cm(1.27)
    section.footer_distance = Cm(1.27)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "My-work-office · 技术验证报告"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "机密文件 · 请勿外传"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for page in range(1, 4):
        doc.add_heading(f"复杂文档 第{page}节", level=1)
        for i in range(1, 4):
            p = doc.add_paragraph(f"[C{page}-{i}] {LOREM}{LOREM}")
            p.paragraph_format.space_after = Pt(6)
        doc.add_heading("附：图片占位段落", level=2)
        pic_p = doc.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pic_p.add_run("[图片占位 — 实际产品将插入 InlineImage]")
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        run.font.italic = True
        for i in range(4, 7):
            p = doc.add_paragraph(f"[C{page}-{i}] {LOREM}")
            p.paragraph_format.space_after = Pt(6)
        if page < 3:
            _add_page_break(doc)
    doc.save(os.path.join(OUT, "complex.docx"))
    print("生成: complex.docx (3页, 页眉页脚+图片占位)")

def make_natural():
    doc = Document()
    doc.core_properties.title = "自然分页测试文档"
    doc.add_heading("自然分页压力测试（无显式分页符）", level=1)
    for i in range(1, 61):
        p = doc.add_paragraph(f"[N1-{i}] {LOREM}{LOREM}")
        p.paragraph_format.space_after = Pt(6)
    doc.save(os.path.join(OUT, "natural.docx"))
    print("生成: natural.docx (自然溢出分页, 60段)")


if __name__ == "__main__":
    make_simple()
    make_medium()
    make_complex()
    make_natural()
    print("全部样本生成完毕 →", OUT)
